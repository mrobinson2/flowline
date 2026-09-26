# -*- coding: utf-8 -*-
"""Builds a 146-row synthetic workbook that matches the real one's shape and totals.
   Task names are invented. Per-phase lead/cycle totals match the Summary by Phase tab
   exactly, so the importer and metrics can be verified against known-good numbers."""
import random, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from schema import *

random.seed(20260920)

# steps that carry a gate / handoff / rework, matching the real per-phase counts
GATES   = {1:5, 2:2, 3:4, 4:2, 5:1, 6:5, 7:1, 8:3, 9:1, 10:0, 11:1, 12:0, 13:1, 14:3, 15:0, 16:2, 17:3, 18:2, 19:0}
HANDOFF = {1:1, 11:1, 17:1, 18:1}
REWORK  = {1:2, 5:1, 13:1, 14:1, 15:1}

VERBS = ["Submitted","Drafted","Reviewed","Approved","Assigned","Validated","Configured","Executed",
         "Documented","Published","Confirmed","Provisioned","Onboarded","Signed off","Completed"]
NOUNS = ["intake record","design package","control evidence","access model","network path","runbook",
         "test plan","cost model","risk assessment","deployment artefact","service record","data mapping",
         "pattern selection","exception request","readiness check"]

rows, edges = [], []
prev_phase_tails = []
anchor = None
branches = []
rid = 0

for pi, (phase, n, lead, cyc, olead, ocyc) in enumerate(PHASES, start=1):
    # split the phase totals across its steps, weighted, then correct the rounding drift
    w = [random.uniform(0.6, 1.6) for _ in range(n)]
    s = sum(w)

    # Current lead and cycle are WHOLE HOURS in the real workbook: every
    # Earliest Finish on its Edges tab lands on an exact eighth of a day, which
    # only happens if each task is a whole number of hours. Distribute by
    # largest remainder so the phase total still lands exactly.
    def whole(total, weights):
        raw = [total * x / sum(weights) for x in weights]
        out = [max(1, int(v)) for v in raw]
        drift = int(round(total - sum(out)))
        order = sorted(range(len(raw)), key=lambda i: raw[i] - int(raw[i]), reverse=(drift > 0))
        i = 0
        while drift != 0 and i < len(order) * 40:
            j = order[i % len(order)]
            if drift > 0:
                out[j] += 1; drift -= 1
            elif out[j] > 1:
                out[j] -= 1; drift += 1
            i += 1
        return out

    leads = whole(int(lead), w)
    cycs  = whole(int(cyc),  w)
    # Optimized values are current x a reduction factor, so decimals are normal
    oleads = [round(olead* x / s, 1) for x in w]
    ocycs  = [round(ocyc * x / s, 1) for x in w]
    for arr, target in ((oleads, olead), (ocycs, ocyc)):
        arr[0] = round(arr[0] + (target - sum(arr)), 1)

    gates, hos, rws = GATES.get(pi,0), HANDOFF.get(pi,0), REWORK.get(pi,0)
    kinds = (["Approval Gate"]*gates + ["Handoff"]*hos + ["Rework Loop"]*rws)
    kinds += ["Value-Add"] * max(0, n - len(kinds))
    random.shuffle(kinds)

    for k in range(n):
        rid += 1
        kind = kinds[k]
        team = random.choice(list(TEAMS))
        tt   = TEAMS[team]
        if kind == "Approval Gate":
            waste, inter = "Waiting", ("Handoff (TT anti-pattern)" if tt.startswith("Governance") else "Collaboration")
        elif kind == "Rework Loop":
            waste, inter = "Defects", "Collaboration"
        elif kind == "Handoff":
            waste, inter = "Transportation", "Handoff (TT anti-pattern)"
        else:
            waste = "-"
            inter = {"Platform":"X-as-a-Service","Enabling":"Facilitating",
                     "Complicated-Subsystem":"X-as-a-Service","Stream-Aligned":"Collaboration"}.get(tt, "Handoff (TT anti-pattern)")
            if random.random() < 0.28:
                kind, waste = "Business Non-Value-Add (Type 1 Muda)", random.choice(["-","Waiting","Extra-Processing","Inventory"])
        lane = "All lanes" if random.random() < 0.62 else random.choice(["Conditional","Standard + Custom","Custom only"])
        when = "Every workload" if lane == "All lanes" else random.choice(
            ["GenAI use cases","Migrations only","Buy or SaaS only","Azure workloads","Tier 0-3",
             "Regulated / personal data","Non-standard decisions","Internet-facing only"])
        # Dependencies with real parallelism. A strictly serial chain would make
        # the critical path equal the sum of all hours, which would hide the
        # single most important distinction in the whole tool: elapsed time is
        # not the sum of the work. So work fans out into concurrent branches and
        # an approval gate joins them back together, which is how these phases
        # actually run - several teams working at once, then a review that waits
        # for all of them.
        if k == 0:
            preds = list(prev_phase_tails)
            branches = [rid]
            anchor = rid
        elif kind == "Approval Gate" and len(branches) > 0:
            preds = list(dict.fromkeys(branches))      # the gate waits for every open branch
            branches = [rid]
            anchor = rid
        elif len(branches) < 4 and random.random() < 0.42:
            preds = [anchor] if anchor else []         # open a new concurrent branch
            branches.append(rid)
        else:
            i = random.randrange(len(branches))        # extend an existing branch
            preds = [branches[i]]
            branches[i] = rid
        for p in preds:
            edges.append((p, rid))
        rows.append([rid, phase, f"{random.choice(NOUNS).capitalize()} {random.choice(VERBS)}".strip(),
                     team, tt, inter, ",".join(str(p) for p in preds), kind, waste, lane, when,
                     leads[k], cycs[k], oleads[k], ocycs[k]])
    prev_phase_tails = list(dict.fromkeys(branches))

# ---------------------------------------------------------------------------
# CPM schedule in days, exactly as the real workbook carries it in columns R..X.
# The app recomputes all of this from the dependencies; having it here means the
# reconciliation path is actually exercised rather than assumed to work.
# ---------------------------------------------------------------------------
NAME = {r[0]: r[2] for r in rows}
PREDS = {r[0]: ([int(x) for x in r[6].split(",")] if r[6] else []) for r in rows}
SUCCS = {r[0]: [] for r in rows}
for p, s in edges:
    SUCCS[p].append(s)

DUR = {r[0]: round((r[11] + r[12]) / HOURS_PER_DAY, 4) for r in rows}
ES, EF = {}, {}
for rid in sorted(DUR):                                   # ids are already topological
    ES[rid] = max([EF[p] for p in PREDS[rid]], default=0.0)
    EF[rid] = round(ES[rid] + DUR[rid], 4)
END = max(EF.values())

LS, LF = {}, {}
for rid in sorted(DUR, reverse=True):
    LF[rid] = min([LS[s] for s in SUCCS[rid]], default=END)
    LS[rid] = round(LF[rid] - DUR[rid], 4)
SLACK = {rid: round(LS[rid] - ES[rid], 4) for rid in DUR}

# %C&A by step type: gates and rework send work back most often
CA = {"Approval Gate": (0.55, 0.80), "Rework Loop": (0.40, 0.65), "Handoff": (0.60, 0.85),
      "Business Non-Value-Add (Type 1 Muda)": (0.70, 0.92), "Value-Add": (0.80, 0.98)}
for r in rows:
    rid, lead, cyc = r[0], r[11], r[12]
    lo, hi = CA.get(r[7], (0.75, 0.95))
    r += [round(random.uniform(lo, hi), 2),                       # P  %C&A
          round(cyc / (lead + cyc), 4) if (lead + cyc) else 0,    # Q  Flow Efficiency
          DUR[rid],                                              # R  Duration (days)
          round(ES[rid], 2), round(EF[rid], 2),                   # S,T
          round(LS[rid], 2), round(LF[rid], 2),                   # U,V
          round(SLACK[rid], 2),                                   # W  Slack
          "Yes" if SLACK[rid] < 0.005 else "",                    # X  Critical Path
          ""]                                                     # Y  Notes

# Edges tab in the real shape: successor first, and BINDING where slack is zero
edge_rows = []
for p, s in edges:
    edge_rows.append([s, p, NAME[s], NAME[p], round(EF[p], 2), round(ES[s], 2),
                      "BINDING" if abs(EF[p] - ES[s]) < 0.005 else ""])

# ---------------------------------------------------------------------------
# TAILORING MATRIX. Assigns every task a scope class from its phase, then
# overlays the modifiers. Deliberately includes one collision (paved road wants
# to skip the architecture governance gates, a brand new cloud service wants
# them back) so the R-beats-N conflict reporting is actually exercised.
# ---------------------------------------------------------------------------
PHASE_NO = {}
for i, (phase, *_rest) in enumerate(PHASES, start=1):
    PHASE_NO[phase] = i

SMALL_WORK = ["workType=Capacity extension", "workType=In-place upgrade", "workType=Decommission"]
SCOPE = {                      # phase number -> scope class
    1: "spine", 17: "spine", 18: "spine", 19: "close",
    2: "discovery", 7: "sourcing", 12: "data", 15: "dr",
    3: "design", 6: "design",
}
# everything not listed is "build"

AI_TASKS = set()               # ~10 tasks become AI-governance-only
for r in rows:
    if PHASE_NO[r[1]] in (1, 5, 6, 8, 14, 17) and r[0] % 13 == 0:
        AI_TASKS.add(r[0])

matrix_rows = []
for r in rows:
    rid, phase, name, kind = r[0], r[1], r[2], r[7]
    pn = PHASE_NO[phase]
    cls = SCOPE.get(pn, "build")
    cell = {c: "" for c in MATRIX_CONDITIONS}
    baseline = "Yes"

    if rid in AI_TASKS:
        baseline, cell["genAI"] = "No", "R"
    elif cls == "discovery":
        baseline, cell["discovery"] = "No", "R"
        if rid % 3 == 0: cell["rfi"] = "R"
        if rid % 3 == 1: cell["rfp"] = "R"
        if rid % 3 == 2: cell["poc"] = "R"
    elif cls == "sourcing":
        baseline = "No"
        cell["newVendor"], cell["hardware"], cell["newLicense"] = "R", "R", "R"
    elif cls == "data":
        baseline, cell["dataMigration"] = "No", "R"
    elif cls == "dr":
        baseline, cell["drRequired"] = "No", "R"
    elif cls == "design":
        for w in SMALL_WORK + ["workType=Lift and shift"]:
            cell[w] = "N"
    elif cls == "build":
        for w in SMALL_WORK:
            cell[w] = "N"
    # "spine" and "close" apply to every work type

    # Overlays below only apply to steps that are in the baseline. Stacking a
    # duration multiplier onto a step that only exists because of a modifier
    # would make that multiplier the sole route in, which is not what a
    # multiplier is for.
    if baseline != "Yes":
        matrix_rows.append([rid, name, baseline] + [cell[c] for c in MATRIX_CONDITIONS])
        continue

    # the paved road retires reviews and approvals it has already passed centrally
    if kind == "Approval Gate" and pn in (5, 6, 9) and baseline == "Yes":
        cell["pavedRoad"] = "N"
        # ...but a service nobody has run here before needs them back, and takes longer
        if pn == 6:
            cell["newService"] = "2.5"
    # a brand new service type makes design work longer wherever it lands
    elif pn == 5 and kind == "Value-Add" and rid % 4 == 0:
        cell["newService"] = "1.8"
    # top-tier services get more scrutiny at the security and readiness gates
    if kind == "Approval Gate" and pn in (8, 16, 17):
        cell["tier=Tier 0|Tier 1"] = "1.5"
    if pn == 8 and kind != "Approval Gate":
        cell["dataClass=Regulated"] = "1.4"
    if pn == 11 and rid % 2 == 0:
        baseline, cell["internetFacing"] = "No", "R"

    matrix_rows.append([rid, name, baseline] + [cell[c] for c in MATRIX_CONDITIONS])

# ---------------------------------------------------------------------------
# RULE-LAYER COLUMNS on the Task List (all optional in the app). The AI
# governance rows carry an Include Expression semantically identical to their
# matrix cell, so the scoped schedule is unchanged and the expression-vs-
# matrix precedence path gets exercised by a real workbook. Phase-8 gates are
# Governed: a normal user cannot switch a security approval off.
# ---------------------------------------------------------------------------
for r in rows:
    rid, phase, kind = r[0], r[1], r[7]
    expr = "genAI = true" if rid in AI_TASKS else ""
    trig = "AI governance applies to generative AI workloads" if rid in AI_TASKS else ""
    can = "Governed" if (kind == "Approval Gate" and PHASE_NO[phase] == 8) else ""
    r += ["", expr, trig, "", can, ""]   # Rule ID, Include Expression, Trigger Explanation, Default Included, Can Override, Rule Priority

wb = openpyxl.Workbook()
hdr_fill = PatternFill("solid", fgColor="FF1F3864")
hdr_font = Font(bold=True, color="FFFFFFFF")

def sheet(name, cols, data, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font = hdr_fill, hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in data:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    for i, wdt in enumerate(widths or [], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
    return ws

sheet("Task List", TASK_COLUMNS + ["Rule ID", "Include Expression", "Trigger Explanation",
                                   "Default Included", "Can Override", "Rule Priority"], rows,
      [6,30,42,26,24,24,14,30,20,16,26,12,12,12,12,9,11,12,12,12,12,12,10,11,20,10,24,34,12,12,11], first=True)
sheet("Edges", EDGE_COLUMNS, edge_rows, [12,13,42,42,14,14,13])
sheet("Toggles", TOGGLE_COLUMNS,
      [t + TOGGLE_EXTRAS.get(t[0], ("", "Always", "", "")) for t in TOGGLES],
      [16, 12, 40, 10, 54, 20, 14, 52, 9, 20, 9, 40])
sheet("Rules", RULE_COLUMNS, RULES, [18, 56, 34, 52])
sheet("Profiles", PROFILE_COLUMNS + [t[0] for t in TOGGLES],
      [[p[0], p[1], p[2]] + [p[3].get(t[0], "") for t in TOGGLES] for p in PROFILES],
      [20, 34, 52] + [13] * len(TOGGLES))
sheet("Scenario Matrix", MATRIX_FIXED + MATRIX_CONDITIONS, matrix_rows,
      [9, 42, 10] + [14] * len(MATRIX_CONDITIONS))
sheet("Assumptions", ASSUMPTION_COLUMNS,
      [["Approval Gate",0.80,0.50],["Handoff",0.85,0.50],["Rework Loop",0.85,0.70],
       ["Business Non-Value-Add (Type 1 Muda)",0.70,0.40],["Waste (Type 2 Muda)",1.00,1.00],
       ["Value-Add",0.60,0.30],["Milestone",0.50,0.00]], [40,26,26])

# Summary by Phase, computed from the rows so it proves the fixture reconciles
summ = []
for phase, n, lead, cyc, olead, ocyc in PHASES:
    pr = [r for r in rows if r[1] == phase]
    L, C = sum(r[11] for r in pr), sum(r[12] for r in pr)
    summ.append([phase, len(pr), sum(1 for r in pr if r[7]=="Approval Gate"),
                 sum(1 for r in pr if r[7]=="Handoff"), sum(1 for r in pr if r[7]=="Rework Loop"),
                 round(L,1), round(C,1), round(sum(r[13] for r in pr),1), round(sum(r[14] for r in pr),1),
                 round(L + C - sum(r[13] for r in pr) - sum(r[14] for r in pr),1),
                 round(C/(L+C)*100,1)])
sheet("Summary by Phase",
      ["Phase","Steps","Approval Gates","Handoffs","Rework Loops","Current Lead (hrs)","Current Cycle (hrs)",
       "Optimized Lead (hrs)","Optimized Cycle (hrs)","Lead Time Reduction (hrs)","Flow Efficiency %"],
      summ, [34,8,14,11,13,16,16,16,16,18,15])

wb.save("sample-value-stream.xlsx")

tl = sum(r[11] for r in rows); tc = sum(r[12] for r in rows)
to = sum(r[13] for r in rows); toc = sum(r[14] for r in rows)
print(f"rows {len(rows)}  edges {len(edges)}")
print(f"lead {tl:.1f} (7992.0)   cycle {tc:.1f} (889.0)")
print(f"optlead {to:.1f} (2243.8)   optcycle {toc:.1f} (523.8)")
print(f"gates {sum(1 for r in rows if r[7]=='Approval Gate')} (36)  "
      f"handoffs {sum(1 for r in rows if r[7]=='Handoff')} (4)  "
      f"rework {sum(1 for r in rows if r[7]=='Rework Loop')} (6)")
print(f"flow efficiency {tc/(tl+tc)*100:.1f}% (10.0)")
print(f"critical path {END:.2f} days  ({END*HOURS_PER_DAY:.0f} hrs)   summed work {sum(DUR.values()):.1f} days")
print(f"binding links {sum(1 for e in edge_rows if e[6]==chr(66)+chr(73)+chr(78)+chr(68)+chr(73)+chr(78)+chr(71))} of {len(edge_rows)}")
print(f"critical tasks {sum(1 for r in rows if r[23]==chr(89)+chr(101)+chr(115))}")
